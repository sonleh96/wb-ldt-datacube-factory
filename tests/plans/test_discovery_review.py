from __future__ import annotations

import csv
import logging
from pathlib import Path

import yaml
from openpyxl import load_workbook

from ldt_factory.plans.config import load_plan_config
from ldt_factory.plans.discovery import discover_area, load_area_selection
from ldt_factory.plans.registry import load_admin_registry
from ldt_factory.plans.review import apply_review_workbook, export_review_workbook
from ldt_factory.plans.scoring import canonicalize_url


def _logger() -> logging.Logger:
    logger = logging.getLogger("test.plans")
    logger.handlers.clear()
    logger.addHandler(logging.NullHandler())
    return logger


def _config(tmp_path: Path):
    registry_path = tmp_path / "admin2.csv"
    with registry_path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=["code", "name", "parent"])
        writer.writeheader()
        writer.writerow({"code": "TST001", "name": "Central", "parent": "North"})
    payload = {
        "schema_version": 1,
        "country": {
            "iso3": "TST",
            "name": "Testland",
            "languages": ["en"],
            "admin2_types": ["district"],
        },
        "workspace": "workflow",
        "admin_registry": {
            "path": registry_path.name,
            "id_field": "code",
            "name_field": "name",
            "parent_field": "parent",
            "alias_fields": [],
        },
        "search": {
            "api_key_env": "EXA_API_KEY",
            "minimum_passes": 2,
            "max_passes": 3,
            "num_results": 10,
            "request_timeout_seconds": 30,
            "document_terms": ["integrated development plan", "IDP"],
            "final_terms": ["approved", "final"],
            "draft_terms": ["draft", "citizen version"],
            "official_domain_suffixes": ["gov.test"],
        },
        "review": {
            "auto_accept_threshold": 0.9,
            "minimum_margin": 0.1,
            "require_formal_adoption_for_auto_accept": True,
            "auto_accept_enabled": False,
        },
        "acquisition": {
            "request_timeout_seconds": 30,
            "download_retries": 2,
            "min_file_bytes": 10,
            "max_file_bytes": 1000000,
        },
        "storage": {"bucket": "bucket", "prefix": "plans/districts"},
    }
    path = tmp_path / "plans.yaml"
    path.write_text(yaml.safe_dump(payload), encoding="utf-8")
    return load_plan_config(path, data_root=tmp_path)


def _result(url: str, *, status: str, full: bool, end_year: str, title: str):
    return {
        "title": title,
        "url": url,
        "publishedDate": "2025-01-01",
        "highlights": [f"Central District Integrated Development Plan 2023-{end_year} {status}"],
        "summary": {
            "jurisdiction": "Central District",
            "jurisdiction_match": True,
            "is_development_plan": True,
            "is_full_document": full,
            "status": status,
            "start_year": "2023",
            "end_year": end_year,
            "adoption_date": "2023-06-01",
            "issuing_authority": "Central District Council",
            "direct_document_url": url,
        },
        "extras": {"links": []},
    }


class FakeExaClient:
    def __init__(self):
        self.requests = []

    def search(self, payload):
        self.requests.append(payload)
        if len(self.requests) == 1:
            results = [
                _result(
                    "https://central.gov.test/draft.pdf?utm_source=search",
                    status="draft",
                    full=True,
                    end_year="2030",
                    title="Draft Central IDP",
                )
            ]
        else:
            results = [
                _result(
                    "https://central.gov.test/approved.pdf",
                    status="approved",
                    full=True,
                    end_year="2033",
                    title="Approved Central IDP",
                )
            ]
        return {"requestId": f"request-{len(self.requests)}", "results": results}


def test_discovery_runs_two_passes_and_selects_approved_official_plan(tmp_path: Path):
    config = _config(tmp_path)
    area = load_admin_registry(config)[0]
    client = FakeExaClient()
    config.prepare_run("run-1")

    selection = discover_area(
        config,
        area,
        run_id="run-1",
        client=client,
        logger=_logger(),
    )

    assert len(client.requests) == 2
    assert client.requests[0]["type"] == "auto"
    assert client.requests[1]["type"] == "deep"
    assert selection.proposed_decision == "AUTO_ACCEPT"
    assert selection.selected is not None
    assert selection.selected.document_url == "https://central.gov.test/approved.pdf"
    assert selection.selected.score == 1.0
    loaded, candidates = load_area_selection(config, "run-1", area)
    assert loaded.proposed_decision == "AUTO_ACCEPT"
    assert len(candidates) == 2


def test_review_workbook_round_trip_records_explicit_approval(tmp_path: Path):
    config = _config(tmp_path)
    area = load_admin_registry(config)[0]
    config.prepare_run("run-1")
    discover_area(config, area, run_id="run-1", client=FakeExaClient(), logger=_logger())
    workbook_path = export_review_workbook(config, "run-1", [area])

    workbook = load_workbook(workbook_path)
    sheet = workbook["Review Queue"]
    headers = [cell.value for cell in sheet[1]]
    sheet.cell(2, headers.index("reviewer_decision") + 1, "APPROVE")
    sheet.cell(2, headers.index("reviewer_notes") + 1, "Checked against council website")
    sheet.cell(2, headers.index("pin_selection") + 1, True)
    workbook.save(workbook_path)
    workbook.close()

    decisions_path = apply_review_workbook(config, "run-1", workbook_path, [area])
    payload = __import__("json").loads(decisions_path.read_text(encoding="utf-8"))
    decision = payload["decisions"][0]
    assert decision["decision"] == "APPROVE"
    assert decision["selected_url"] == "https://central.gov.test/approved.pdf"
    assert decision["pinned"] is True


def test_url_canonicalization_removes_tracking_and_fragments():
    assert canonicalize_url("HTTPS://Example.COM/a.pdf?utm_source=x&id=2#page=3") == (
        "https://example.com/a.pdf?id=2"
    )
