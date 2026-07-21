from __future__ import annotations

import datetime as dt
import json
from collections import Counter
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from ..domains._api_cache import atomic_write_json
from .config import DevelopmentPlanConfig
from .discovery import load_area_selection
from .models import AreaSelection, PlanCandidate
from .registry import AdminArea
from .scoring import canonicalize_url


REVIEW_COLUMNS = (
    "admin2_id",
    "admin2_name",
    "parent",
    "storage_slug",
    "proposed_decision",
    "score",
    "title",
    "candidate_url",
    "document_url",
    "source_tier",
    "document_status",
    "document_type",
    "start_year",
    "end_year",
    "adoption_date",
    "issuing_authority",
    "evidence",
    "reason_codes",
    "reviewer_decision",
    "replacement_url",
    "reviewer_notes",
    "pin_selection",
)


def _selection_row(selection: AreaSelection) -> list[Any]:
    candidate = selection.selected
    return [
        selection.area.admin2_id,
        selection.area.name,
        selection.area.parent,
        selection.area.storage_slug,
        selection.proposed_decision,
        candidate.score if candidate else None,
        candidate.title if candidate else "",
        candidate.landing_url if candidate else "",
        candidate.document_url if candidate else "",
        candidate.source_tier if candidate else "",
        candidate.document_status if candidate else "",
        candidate.document_type if candidate else "",
        candidate.start_year if candidate else None,
        candidate.end_year if candidate else None,
        candidate.adoption_date if candidate else "",
        candidate.issuing_authority if candidate else "",
        "\n".join(candidate.evidence) if candidate else "",
        ";".join(selection.reason_codes),
        "",
        "",
        "",
        False,
    ]


def _style_table(sheet, *, widths: dict[str, int] | None = None) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        letter = column[0].column_letter
        heading = str(column[0].value or "")
        sheet.column_dimensions[letter].width = (widths or {}).get(heading, 18)
        for cell in column[1:]:
            cell.alignment = Alignment(wrap_text=True, vertical="top")


def export_review_workbook(
    config: DevelopmentPlanConfig,
    run_id: str,
    areas: list[AdminArea],
    *,
    output: Path | None = None,
) -> Path:
    selections: list[AreaSelection] = []
    candidates: list[PlanCandidate] = []
    for area in areas:
        try:
            selection, area_candidates = load_area_selection(config, run_id, area)
        except FileNotFoundError:
            selection = AreaSelection(area, "MISSING", None, ("not_searched",), 0)
            area_candidates = []
        selections.append(selection)
        candidates.extend(area_candidates)

    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    counts = Counter(selection.proposed_decision for selection in selections)
    summary.append(["Status", "Count"])
    for status in ("AUTO_ACCEPT", "REVIEW", "MISSING"):
        summary.append([status, counts[status]])
    summary.append(["TOTAL", len(selections)])
    _style_table(summary, widths={"Status": 24, "Count": 14})

    review = workbook.create_sheet("Review Queue")
    review.append(REVIEW_COLUMNS)
    for selection in selections:
        review.append(_selection_row(selection))
    _style_table(
        review,
        widths={
            "admin2_name": 24,
            "parent": 20,
            "title": 42,
            "candidate_url": 45,
            "document_url": 45,
            "evidence": 70,
            "reason_codes": 36,
            "replacement_url": 45,
            "reviewer_notes": 40,
        },
    )
    decision_column = REVIEW_COLUMNS.index("reviewer_decision") + 1
    replacement_column = REVIEW_COLUMNS.index("replacement_url") + 1
    notes_column = REVIEW_COLUMNS.index("reviewer_notes") + 1
    pin_column = REVIEW_COLUMNS.index("pin_selection") + 1
    editable_fill = PatternFill("solid", fgColor="FFF2CC")
    for row in range(2, review.max_row + 1):
        for column in (decision_column, replacement_column, notes_column, pin_column):
            review.cell(row, column).fill = editable_fill
    decision_validation = DataValidation(type="list", formula1='"APPROVE,REJECT,MISSING"')
    pin_validation = DataValidation(type="list", formula1='"TRUE,FALSE"')
    review.add_data_validation(decision_validation)
    review.add_data_validation(pin_validation)
    decision_validation.add(f"{review.cell(2, decision_column).coordinate}:{review.cell(review.max_row, decision_column).coordinate}")
    pin_validation.add(f"{review.cell(2, pin_column).coordinate}:{review.cell(review.max_row, pin_column).coordinate}")
    review.conditional_formatting.add(
        f"A2:V{review.max_row}",
        FormulaRule(formula=[f'${review.cell(2, decision_column).column_letter}2="APPROVE"'], fill=PatternFill("solid", fgColor="E2F0D9")),
    )

    all_candidates = workbook.create_sheet("All Candidates")
    candidate_columns = tuple(PlanCandidate.__dataclass_fields__)
    all_candidates.append(candidate_columns)
    for candidate in candidates:
        row = candidate.to_dict()
        all_candidates.append(
            [
                "\n".join(value) if isinstance(value := row[column], list) else value
                for column in candidate_columns
            ]
        )
    _style_table(
        all_candidates,
        widths={"title": 42, "landing_url": 45, "document_url": 45, "query": 60, "evidence": 70},
    )

    metadata = workbook.create_sheet("Run Metadata")
    metadata.append(["Field", "Value"])
    for key, value in (
        ("run_id", run_id),
        ("country", f"{config.country_name} ({config.iso3})"),
        ("config", str(config.path)),
        ("registry", str(config.registry_path)),
        ("generated_at", dt.datetime.now(dt.timezone.utc).isoformat()),
        ("candidate_count", len(candidates)),
        ("instructions", "Edit only the yellow reviewer columns. Replacement URLs are revalidated before acquisition."),
    ):
        metadata.append([key, value])
    _style_table(metadata, widths={"Field": 24, "Value": 100})

    destination = output or config.run_dir(run_id) / "review.xlsx"
    destination.parent.mkdir(parents=True, exist_ok=True)
    temporary = destination.with_suffix(destination.suffix + ".part")
    workbook.save(temporary)
    temporary.replace(destination)
    return destination


def apply_review_workbook(
    config: DevelopmentPlanConfig,
    run_id: str,
    workbook_path: Path,
    areas: list[AdminArea],
) -> Path:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    if "Review Queue" not in workbook.sheetnames:
        raise ValueError("Review workbook has no 'Review Queue' sheet")
    sheet = workbook["Review Queue"]
    headers = [str(cell.value or "").strip() for cell in sheet[1]]
    missing = sorted(set(REVIEW_COLUMNS) - set(headers))
    if missing:
        raise ValueError(f"Review Queue is missing columns: {', '.join(missing)}")
    indexes = {name: headers.index(name) for name in REVIEW_COLUMNS}
    area_by_id = {area.admin2_id: area for area in areas}
    decisions: list[dict[str, Any]] = []
    seen: set[str] = set()
    for row_number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        admin2_id = str(values[indexes["admin2_id"]] or "").strip()
        if not admin2_id:
            continue
        if admin2_id not in area_by_id:
            raise ValueError(f"Unknown admin2_id {admin2_id!r} on row {row_number}")
        if admin2_id in seen:
            raise ValueError(f"Duplicate admin2_id {admin2_id!r} on row {row_number}")
        seen.add(admin2_id)
        decision = str(values[indexes["reviewer_decision"]] or "").strip().upper()
        if not decision:
            continue
        if decision not in {"APPROVE", "REJECT", "MISSING"}:
            raise ValueError(f"Invalid reviewer_decision {decision!r} on row {row_number}")
        replacement = canonicalize_url(str(values[indexes["replacement_url"]] or ""))
        candidate_url = canonicalize_url(str(values[indexes["document_url"]] or "")) or canonicalize_url(
            str(values[indexes["candidate_url"]] or "")
        )
        selected_url = replacement or candidate_url
        if decision == "APPROVE" and not selected_url:
            raise ValueError(f"APPROVE row {row_number} has no usable candidate or replacement URL")
        pin_raw = values[indexes["pin_selection"]]
        pinned = pin_raw is True or str(pin_raw).strip().upper() == "TRUE"
        area = area_by_id[admin2_id]
        decisions.append(
            {
                "admin2_id": admin2_id,
                "admin2_name": area.name,
                "parent": area.parent,
                "storage_slug": area.storage_slug,
                "decision": decision,
                "selected_url": selected_url,
                "replacement_url": replacement,
                "reviewer_notes": str(values[indexes["reviewer_notes"]] or "").strip(),
                "pinned": pinned,
                "reviewed_at": dt.datetime.now(dt.timezone.utc).isoformat(),
            }
        )
    workbook.close()
    payload = {
        "schema_version": 1,
        "run_id": run_id,
        "country": config.iso3,
        "source_workbook": str(workbook_path.resolve()),
        "decisions": decisions,
    }
    destination = config.run_dir(run_id) / "decisions" / "reviewed.json"
    atomic_write_json(destination, payload)
    return destination
